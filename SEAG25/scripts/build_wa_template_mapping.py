#!/usr/bin/env python3
"""
Generate the WA Template Mapping sheet for SEAG25.

This script reads the `(EDIT HERE) sport_event_template_mapping.xlsx`
workbook, normalises the WhatsApp templates (splitting individual/team
and final/non-final variants), and pushes the results to the
`WA Template Mapping` sheet in the master Google Sheet.
"""
import copy
import json
import re
from collections import defaultdict
from pathlib import Path

import gspread
import openpyxl
from google.oauth2.service_account import Credentials

SPREADSHEET_ID = "15zXDQdkGeAN2AMMdrJ_qjkjReq_Y3mlU1-_7ciniyS4"
WA_SHEET_NAME = "WA Template Mapping"
SCHEDULE_SHEET_NAME = "SEAG25 Competition Schedule"
TEMPLATE_FILE = Path("SEAG25/(EDIT HERE) sport_event_template_mapping.xlsx")
TEMPLATE_TAB = "Summarised WA templates"

SERVICE_ACCOUNT_FILE = Path(
    "AYG25/ayg-form-system/functions/google_credentials.json"
)

STATUS_PLACEHOLDER = "{ADVANCEMENT_STATUS} {MEDAL_STATUS}"
PB_PLACEHOLDER = "{PB_NR_TEXT}"
PB_ALIASES = [
    "{PB/NR/GR}",
    "{PB/NR}",
    "{PB}",
    "{PB/NR/GR.}",
]
ADVANCEMENT_ALIASES = [
    "{ADVANCEMENT / MEDAL_STATUS}",
    "{ADVANCEMENT/ MEDAL_STATUS}",
    "{ADVANCEMENT / MEDAL STATUS}",
]

PLACEHOLDER_PATTERN = re.compile(r"{[A-Z0-9_ ]+}")
HEADER_PATTERN = re.compile(r"(\*{SPORT} - [^\n]+{ROUND}\*)", re.IGNORECASE)
MEDAL_ROUND_KEYWORDS = ["BRONZE", "GOLD", "SILVER", "MEDAL", "FINAL"]
SIMPLE_PLACEHOLDER_ALIASES = {
    "{ATHLETE1}": "{NAME}",
    "{ATHLETE2}": "{COMPETITOR}",
    "{ATHLETE2_COUNTRY}": "{COUNTRY2}",
    "{ATHLETE2_NAMES}": "{COMPETITOR}",
    "{OPPONENT_COUNTRY}": "{COUNTRY2}",
    "{TOTAL_PARTICIPANTS}": "{TOTAL}",
    "{CURRENT ROUND}": "{ROUND}",
    "{ROUND_NUMBER}": "{ROUND}",
    "{RANK}": "{PLACEMENT}",
    "{RESULT}": "{SCORE}",
    "{TOTAL_SCORE_INFO}": "{WIN_TYPE}",
    "{PAR_STATUS}": "{WIN_TYPE}",
    "{ADVANCEMENT/MEDAL_STATUS}": STATUS_PLACEHOLDER,
    "{NAME(s)}": "{NAMES}",
    "{Names}": "{NAMES}",
    "{NAME(S)}": "{NAMES}",
}

PAIR_PLACEHOLDER_PATTERNS = [
    (re.compile(r"{ATHLETE A}\s*,\s*{ATHLETE B}", re.IGNORECASE), "{NAMES}"),
    (re.compile(r"{ATHLETE A}", re.IGNORECASE), "{NAMES}"),
    (re.compile(r"{ATHLETE B}", re.IGNORECASE), "{NAMES}"),
]

SHOOTING_PISTOL_PRELIM_1 = """*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*

{NAME} (SGP)

Score: {SCORE}.

Currently ranked {PLACEMENT} out of {TOTAL} after {ROUND} {MEDAL_STATUS}

{PB_NR_TEXT}"""

SHOOTING_PISTOL_PRELIM_2 = """*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*

{NAME} (SGP)

Score: xx

Total Score: {SCORE}.

SGP finished {PLACEMENT} out of {TOTAL} {MEDAL_STATUS}

{PB_NR_TEXT}"""

SHOOTING_INDIV_TEMPLATE = """*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*

{NAME} (SGP)

Score: {SCORE}.

SGP finished {PLACEMENT} out of {TOTAL}. {ADVANCEMENT_STATUS} {MEDAL_STATUS}

{PB_NR_TEXT}"""

SHOOTING_TEAM_TEMPLATE = """*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*

Team of {NAMES} (SGP)

Score: {SCORE}.

Finished {PLACEMENT} out of {TOTAL}. {ADVANCEMENT_STATUS} {MEDAL_STATUS}

{PB_NR_TEXT}"""

SHOOTING_TEAM_FINAL_TEMPLATE = """*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*

Team of {NAMES} (SGP)

SGP finished {PLACEMENT} out of {TOTAL}. {ADVANCEMENT_STATUS} {MEDAL_STATUS}

{PB_NR_TEXT}"""


def safe_upper(value: str) -> str:
    return value.replace("\n", " ").strip().upper() if isinstance(value, str) else ""


def load_schedule_rows(client):
    ws = client.open_by_key(SPREADSHEET_ID).worksheet(SCHEDULE_SHEET_NAME)
    values = ws.get_all_values()
    headers = values[7]
    index = {name: idx for idx, name in enumerate(headers)}
    return values[8:], index


def collect_events(rows, idx, filter_fn):
    sports, disciplines, events, rounds = set(), set(), set(), set()
    for row in rows:
        sport = row[idx["SPORT"]]
        discipline = row[idx["DISCIPLINE"]]
        event = row[idx["EVENT"]]
        round_name = row[idx["STAGE / ROUND OF COMPETITION"]]
        team_flag = row[idx["TEAM SPORT/EVENT"]]

        if not sport:
            continue
        if safe_upper(event) == "TECHNICAL MEETING":
            continue
        if not filter_fn(row, sport, discipline, event, round_name, team_flag):
            continue
        sports.add(safe_upper(sport))
        disciplines.add(safe_upper(discipline))
        events.add(safe_upper(event))
        if round_name:
            rounds.add(safe_upper(round_name))

    return sports, disciplines, events, rounds


def normalize_template_text(text: str) -> str:
    if not text:
        return ""
    for alias in ADVANCEMENT_ALIASES:
        text = text.replace(alias, STATUS_PLACEHOLDER)
    for alias in PB_ALIASES:
        text = text.replace(alias, PB_PLACEHOLDER)
    text = text.replace("{THA TIME}", "{TIME_THA}")
    if PB_PLACEHOLDER not in text:
        text = text.rstrip() + "\n\n{PB_NR_TEXT}"
    text = normalize_placeholder_aliases(text)
    return format_template_layout(text)


def format_template_layout(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    def header_repl(match):
        return match.group(1).strip() + "\n\n"

    text = HEADER_PATTERN.sub(header_repl, text, count=1)
    # Ensure blank line after "(SGP)" or "Team of ..." lines before Time/Score
    text = re.sub(r"(\(SGP\))\s+(Time:)", r"\1\n\n\2", text)
    text = re.sub(r"(\(SGP\))\s+(Score:)", r"\1\n\n\2", text)
    text = re.sub(r"(SGP VS {COUNTRY})\s+", r"\1\n\n", text)
    text = re.sub(r"(Score:\s*{SCORE}\.)\s+", r"\1\n\n", text)
    text = re.sub(r"{NAMES}\s*\(SGP\)", "{NAMES} (SGP)", text, flags=re.IGNORECASE)
    text = re.sub(r"{NAME}\s*\(SGP\)", "{NAME} (SGP)", text, flags=re.IGNORECASE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = text.replace("  ", " ")
    text = text.replace("..", ".")
    text = text.replace("{PB_NR_TEXT}.", "{PB_NR_TEXT}")
    text = re.sub(r"\.{2,}", ".", text)
    text = re.sub(r"\n{0,1}{PB_NR_TEXT}", "\n{PB_NR_TEXT}", text)
    return text.strip()


def split_team_variants(entry):
    """Split templates that contain an 'OR Team of {NAMES}' section."""
    text = entry["template_raw"]
    base_patterns = [
        re.compile(
            r"({NAME}\s*\(SGP\))\s*(?:\n\s*)*OR\s*(?:\n\s*)*(Team of {NAMES}\s*\(SGP\))",
            re.IGNORECASE,
        ),
        re.compile(
            r"(Team of {NAMES}\s*\(SGP\))\s*(?:\n\s*)*OR\s*(?:\n\s*)*({NAME}\s*\(SGP\))",
            re.IGNORECASE,
        ),
    ]
    for pattern in base_patterns:
        match = pattern.search(text)
        if match:
            prefix = text[: match.start()]
            suffix = text[match.end() :]
            first, second = match.group(1), match.group(2)

            # Determine which block is the individual vs team
            if "Team of" in first:
                team_block, indiv_block = first, second
            else:
                indiv_block, team_block = first, second

            indiv_entry = entry.copy()
            indiv_entry["template_raw"] = prefix + indiv_block + suffix
            indiv_entry["team_type"] = "indiv"

            team_entry = entry.copy()
            team_entry["template_raw"] = prefix + team_block + suffix
            team_entry["team_type"] = "team"

            return [indiv_entry, team_entry]

    return [entry]


def infer_team_type_from_text(text: str, default: str) -> str:
    upper = safe_upper(text)
    if "SGP VS" in upper or "TEAM OF" in upper or "TEAM " in upper:
        return "team"
    if "{NAMES}" in upper and "{NAME}" not in upper:
        return "team"
    if "{NAME}" in upper and "TEAM OF" not in upper and "SGP VS" not in upper:
        return "indiv"
    return default


def split_or_blocks(entry):
    text = entry["template_raw"]
    sections = re.split(r"\n\s*OR\s*\n", text)
    if len(sections) <= 1:
        return [entry]

    result = []
    for section in sections:
        cleaned = section.strip()
        if "*{SPORT}" in cleaned and not cleaned.startswith("*{SPORT}"):
            cleaned = cleaned[cleaned.index("*{SPORT}") :]
        new_entry = entry.copy()
        new_entry["template_raw"] = cleaned
        new_entry["team_type"] = infer_team_type_from_text(cleaned, entry["team_type"])
        result.append(new_entry)
    return result


def strip_placeholder(text: str, placeholder: str) -> str:
    text = text.replace(placeholder, "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r" {2,}", " ", text)
    text = text.replace(" .", ".")
    return text.strip()


def normalize_placeholder_aliases(text: str) -> str:
    for pattern, replacement in PAIR_PLACEHOLDER_PATTERNS:
        text = pattern.sub(replacement, text)
    for old, new in SIMPLE_PLACEHOLDER_ALIASES.items():
        text = text.replace(old, new)
    text = text.replace("Team of {Names}", "Team of {NAMES}")
    text = text.replace("Team of {NAME}", "Team of {NAMES}")
    text = text.replace("{NAME(s)}", "{NAMES}")
    return text


def normalize_token(value: str) -> str:
    return re.sub(r'[^A-Z0-9]', '', safe_upper(value))


def text_matches(value: str, target: str) -> bool:
    if not target:
        return True
    value_norm = normalize_token(value)
    target_norm = normalize_token(target)
    if value_norm == target_norm:
        return True
    if value_norm.endswith('S') and value_norm[:-1] == target_norm:
        return True
    if target_norm.endswith('S') and target_norm[:-1] == value_norm:
        return True
    return False


def normalize_team_type(raw: str) -> str:
    if not raw:
        return "all"
    cleaned = raw.lower().replace(" ", "")
    if cleaned in ("all", "both", "any"):
        return "all"
    has_team = "team" in cleaned or "duet" in cleaned
    has_indiv = "indiv" in cleaned or "solo" in cleaned
    if has_team and has_indiv:
        return "all"
    if has_team:
        return "team"
    return "indiv"

def normalize_team_flag(raw: str) -> str:
    if not raw:
        return "indiv"
    cleaned = raw.strip().lower()
    if cleaned in ("1", "team", "teams"):
        return "team"
    if cleaned in ("0", "indiv", "individual", "solo"):
        return "indiv"
    return "indiv"


def has_real_values(value: str) -> bool:
    if not value:
        return False
    stripped = value.strip()
    if not stripped or stripped in ("—", "-"):
        return False
    if stripped.upper() == "ANY":
        return False
    return True


def expand_athletics_rows(row_data):
    sport = safe_upper(row_data[1])
    template = row_data[8]
    if sport != "ATHLETICS":
        return [row_data]
    result_type = (row_data[7] or "").lower()
    if "TIME/DISTANCE/HEIGHT" not in template.upper() and "time, distance, height" not in result_type:
        return [row_data]

    expanded = []
    is_team = (row_data[6] or "").lower() == "team"
    events_str = (row_data[3] or "").upper()
    if is_team and "RELAY" in events_str:
        replacements = [("Time", "TIME")]
    else:
        replacements = [
            ("Time", "TIME"),
            ("Distance", "DISTANCE"),
            ("Height", "HEIGHT"),
        ]
    team_line = (row_data[6] or "").lower() == "team"
    name_line = "Team of {NAMES} (SGP)" if team_line else "{NAME} (SGP)"
    header = "*{SPORT} - {GENDER} {EVENT_NAME} {ROUND}*"

    for label, suffix in replacements:
        new_row = copy.deepcopy(row_data)
        new_row[0] = f"{new_row[0]}_{suffix}"
        new_row[8] = (
            f"{header}\n\n"
            f"{name_line}\n\n"
            f"{label}: {{SCORE}}. Finished {{PLACEMENT}} out of {{TOTAL}}. "
            f"{{ADVANCEMENT_STATUS}} {{MEDAL_STATUS}}\n\n"
            "{PB_NR_TEXT}"
        )
        expanded.append(new_row)
    return expanded


def build_filter(entry):
    target_sport = safe_upper(entry["sport"])
    target_discipline = safe_upper(entry["discipline"])
    discipline_required = bool(target_discipline and target_discipline != target_sport)

    team_type = entry["team_type"]

    def matches(row, sport, discipline, event, round_name, team_flag):
        if not text_matches(sport, target_sport):
            return False
        if discipline_required and not text_matches(discipline, target_discipline):
            return False
        round_upper = safe_upper(round_name)
        if round_upper in ("", "TECHNICAL MEETING"):
            return False

        tf = safe_upper(team_flag)
        if tf in ("0", "1"):
            is_team_row = tf == "1"
        elif tf in ("TEAM", "TEAMS"):
            is_team_row = True
        elif tf in ("INDIVIDUAL", "INDIV"):
            is_team_row = False
        else:
            is_team_row = None

        if is_team_row is not None:
            if team_type == "team" and not is_team_row:
                return False
            if team_type == "indiv" and is_team_row:
                return False

        return True

    return matches


def is_medal_round(round_name: str) -> bool:
    upper = safe_upper(round_name)
    if "SEMI" in upper:
        return False
    return any(keyword in upper for keyword in MEDAL_ROUND_KEYWORDS)


def split_round_variants(template_text, rounds, is_h2h):
    has_medal = "{MEDAL_STATUS}" in template_text
    has_adv = "{ADVANCEMENT_STATUS}" in template_text

    medal_rounds = []
    non_medal_rounds = []
    semi_h2h_rounds = []

    for r in rounds:
        upper = safe_upper(r)
        if is_h2h and "SEMI" in upper:
            semi_h2h_rounds.append(r)
        elif is_medal_round(r):
            medal_rounds.append(r)
        else:
            non_medal_rounds.append(r)

    medal_rounds.sort()
    non_medal_rounds.sort()
    semi_h2h_rounds.sort()

    variants = []
    if medal_rounds and has_medal:
        variants.append(
            (
                "medal",
                medal_rounds,
                strip_placeholder(template_text, "{ADVANCEMENT_STATUS}"),
            )
        )
    if non_medal_rounds and has_adv:
        variants.append(
            (
                "nonmedal",
                non_medal_rounds,
                strip_placeholder(template_text, "{MEDAL_STATUS}"),
            )
        )
    if semi_h2h_rounds:
        variants.append(
            (
                "semi",
                semi_h2h_rounds,
                template_text,
            )
        )

    if not variants:
        base_rounds = sorted(rounds) if rounds else ["ANY"]
        if has_medal and not has_adv:
            variants.append(("medal", base_rounds, template_text))
        elif has_adv and not has_medal:
            variants.append(("nonmedal", base_rounds, template_text))
        else:
            variants.append(("all", base_rounds, template_text))

    return variants


def select_shooting_template(event, round_name, team_flag):
    event_upper = safe_upper(event)
    round_upper = safe_upper(round_name)
    team_type = normalize_team_flag(team_flag)

    if "25M PISTOL" in event_upper:
        if "PRELIMINARY 1" in round_upper:
            return SHOOTING_PISTOL_PRELIM_1, "indiv", False
        if "PRELIMINARY 2" in round_upper:
            return SHOOTING_PISTOL_PRELIM_2, "indiv", False

    if team_type == "team":
        template = SHOOTING_TEAM_FINAL_TEMPLATE if is_medal_round(round_upper) else SHOOTING_TEAM_TEMPLATE
        return template, "team", True

    return SHOOTING_INDIV_TEMPLATE, "indiv", False


def generate_shooting_rows(schedule_rows, idx):
    combos = {}
    for row in schedule_rows:
        if safe_upper(row[idx["SPORT"]]) != "SHOOTING":
            continue
        event = row[idx["EVENT"]] or ""
        round_name = row[idx["STAGE / ROUND OF COMPETITION"]] or ""
        if safe_upper(event) == "TECHNICAL MEETING" or safe_upper(round_name) == "TECHNICAL MEETING":
            continue
        team_flag = row[idx["TEAM SPORT/EVENT"]] or ""
        key = (event, round_name, team_flag)
        group = combos.setdefault(
            key,
            {"sports": set(), "disciplines": set(), "events": set(), "rounds": set()},
        )
        group["sports"].add(row[idx["SPORT"]])
        group["disciplines"].add(row[idx["DISCIPLINE"]])
        if event:
            group["events"].add(event)
        if round_name:
            group["rounds"].add(round_name)

    shooting_rows = []
    for (event, round_name, team_flag), info in combos.items():
        template, team_type, is_h2h = select_shooting_template(event, round_name, team_flag)
        if not template:
            continue
        key = f"SHOOTING_{normalize_token(event) or 'EVENT'}_{normalize_token(round_name) or 'ANY'}_{team_type.upper()}"
        placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(template)))
        field_map = {ph: "computed" for ph in placeholders}
        shooting_rows.append([
            key,
            "\n".join(sorted(info["sports"])) or "SHOOTING",
            "\n".join(sorted(info["disciplines"])) or "SHOOTING",
            "\n".join(sorted(info["events"])) or (event or "—"),
            "\n".join(sorted(info["rounds"])) or (round_name or "—"),
            "yes" if is_h2h else "no",
            team_type,
            "score",
            template,
            json.dumps(placeholders, ensure_ascii=False),
            json.dumps(field_map, ensure_ascii=False),
            "",
        ])
    return shooting_rows


def main():
    creds_info = json.loads(SERVICE_ACCOUNT_FILE.read_text())
    creds = Credentials.from_service_account_info(
        creds_info,
        scopes=[
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ],
    )
    client = gspread.authorize(creds)
    schedule_rows, schedule_idx = load_schedule_rows(client)

    wb = openpyxl.load_workbook(TEMPLATE_FILE, data_only=True)
    sheet = wb[TEMPLATE_TAB]

    entries = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        sport = row[1]
        if not sport:
            continue
        if safe_upper(sport) == "SHOOTING":
            continue

        entry = {
            "sport": sport.strip(),
            "discipline": (row[2] or "").strip(),
            "h2h": (row[3] or "").strip().lower(),
            "team_type": normalize_team_type(row[4]),
            "result_type": (row[5] or "").strip().lower(),
            "template_raw": normalize_template_text(row[6] or ""),
            "sample": row[7] or "",
        }

        for split_entry in split_or_blocks(entry):
            entries.extend(split_team_variants(split_entry))

    sheet_rows = []
    for idx_entry, entry in enumerate(entries, start=1):
        filter_fn = build_filter(entry)
        sports, disciplines, events, rounds = collect_events(
            schedule_rows,
            schedule_idx,
            filter_fn,
        )

        round_variants = split_round_variants(
            entry["template_raw"],
            rounds or set(),
            entry["h2h"] == "yes",
        )
        for suffix, round_subset, template_text in round_variants:
            placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(template_text)))
            field_map = {ph: "computed" for ph in placeholders}
            row_data = [
                f"{safe_upper(entry['sport']).replace(' ', '_')}_{idx_entry:03d}_{suffix.upper()}",
                "\n".join(sorted(sports)) or safe_upper(entry["sport"]),
                "\n".join(sorted(disciplines)) or safe_upper(entry["discipline"]),
                "\n".join(sorted(events)) or "—",
                "\n".join(round_subset) or "—",
                entry["h2h"],
                entry["team_type"],
                entry["result_type"],
                template_text,
                json.dumps(placeholders, ensure_ascii=False),
                json.dumps(field_map, ensure_ascii=False),
                entry["sample"],
              ]
            sheet_rows.extend(expand_athletics_rows(row_data))
    sheet_rows.extend(generate_shooting_rows(schedule_rows, schedule_idx))
    sheet_rows = [
        row for row in sheet_rows
        if has_real_values(row[3]) and has_real_values(row[4])
    ]

    ws = client.open_by_key(SPREADSHEET_ID).worksheet(WA_SHEET_NAME)
    headers = [
        "Template Key",
        "Sports",
        "Disciplines",
        "Events",
        "Rounds",
        "H2H",
        "Team Type",
        "Result Type",
        "Template",
        "Placeholders",
        "Field Map JSON",
        "Sample",
    ]
    ws.clear()
    ws.update("A1", [headers])
    if sheet_rows:
        ws.update("A2", sheet_rows)

    print(f"Wrote {len(sheet_rows)} template rows to '{WA_SHEET_NAME}'.")


if __name__ == "__main__":
    main()
